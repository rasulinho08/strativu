#!/usr/bin/env python3
"""Join azlib/*.json into the sample AZ obligations library (xlsx) and a GRC360 import CSV."""
import csv
import json
import pathlib
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

def g(d, *keys):
    for k in keys:
        v = d.get(k)
        if v:
            if isinstance(v, dict):
                return "; ".join(f"{a}: {b}" for a, b in v.items() if b)
            if isinstance(v, list):
                return "; ".join(x if isinstance(x, str) else json.dumps(x, ensure_ascii=False) for x in v)
            return str(v)
    return ""


SRC = pathlib.Path(__file__).parent  # folder with the *.json files
OUT = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else SRC / "out")
OUT.mkdir(parents=True, exist_ok=True)

FILES = ["personal_data.json", "info_esign_kii.json", "cbar.json"]
laws, obls, pols, notes = [], [], [], []
for f in FILES:
    p = SRC / f
    if not p.exists():
        print("missing", f)
        continue
    d = json.loads(p.read_text(encoding="utf-8"))
    ls = d.get("laws") or ([d["law"]] if d.get("law") else [])
    laws += ls
    obls += d.get("obligations", [])
    for x in d.get("policies", []):
        x["_from"] = g(x, "law", "law_short", "source") or ", ".join(g(l, "law_short", "short", "short_name", "title_az", "title") for l in ls)
        pols.append(x)
    if d.get("notes"):
        notes.append((f, d["notes"] if isinstance(d["notes"], str) else json.dumps(d["notes"], ensure_ascii=False)))

CAT = {"PERSONAL_DATA": "Fərdi məlumatlar", "KII_SECURITY": "KII təhlükəsizliyi", "CENTRAL_BANK": "Mərkəzi Bank",
       "E_SIGNATURE": "Elektron imza", "STATE_STANDARDS": "Dövlət standartları"}
RISK = {"HIGH": "Yüksək", "MEDIUM": "Orta", "LOW": "Aşağı"}

F = "Arial"
HEAD = PatternFill("solid", fgColor="0B3A4A")
HFONT = Font(name=F, bold=True, color="FFFFFF", size=10)
BODY = Font(name=F, size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
thin = Side(style="thin", color="D5DBE3")
RISKFILL = {"HIGH": "FDE2E1", "MEDIUM": "FFF1D6", "LOW": "E3F4E8"}


def j(v):
    if isinstance(v, list):
        return "\n".join(f"• {x}" for x in v)
    return "" if v is None else str(v)


def sheet(wb, title, headers, rows, widths, table_name, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font, c.fill, c.alignment = HFONT, HEAD, Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.alignment, c.border = BODY, WRAP, Border(bottom=thin)
    ws.freeze_panes = "A2"
    if rows:
        t = Table(displayName=table_name, ref=f"A1:{get_column_letter(len(headers))}{len(rows) + 1}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
        ws.add_table(t)
    return ws


wb = Workbook()
# 1. Read me
ws = wb.active
ws.title = "Oxu"
lines = [
    ("GRC360 — Azərbaycan qanunları üzrə nümunə öhdəlik kitabxanası", True),
    ("Bu fayl nümunədir: hər öhdəlik qanunun rəsmi mətnindən sitatla (verbatim) və maddə nömrəsi ilə çıxarılıb. "
     "İstifadədən əvvəl hüquqşünas tərəfindən yoxlanmalıdır; bu hüquqi məsləhət deyil.", False),
    ("", False),
    ("Vərəqlər", True),
    ("Mənbələr — qanun və qaydalar, rəsmi nömrə, tarix, nəzarət orqanı və mətnin götürüldüyü link.", False),
    ("Öhdəliklər — hər sətir bir öhdəlik: maddə, rəsmi mətndən sitat, sadə dildə öhdəlik (AZ/EN), kimə aiddir, tezlik/son tarix, risk, təklif olunan control-lar və policy-lər.", False),
    ("Policy-lər — bu qanunlara görə şirkətdə olmalı daxili siyasət və prosedurlar, hansı maddələrdən irəli gəlir.", False),
    ("Xülasə — qanun və risk səviyyəsinə görə saylar (düsturla hesablanır).", False),
    ("Qeydlər — mətnləri çıxararkən qarşılaşılan məhdudiyyətlər və yoxlanmalı məqamlar.", False),
    ("", False),
    ("GRC360-a import", True),
    ("Yanındakı GRC360-obligations-import.csv faylı Compliance Hub › Obligations › Import CSV ilə birbaşa yüklənir "
     "(sütunlar: legalActName, articleOrClauseNo, category, requirementDescription, regulatoryAuthority, status, riskLevel). "
     "Bütün öhdəliklər 'Under Review' statusu ilə gəlir; Responsible Team import-dan sonra təyin edilir.", False),
    ("", False),
    ("Sütunlar haqqında", True),
    ("Risk — pozuntunun mümkün nəticəsinə görə ilkin qiymətləndirmədir (Yüksək / Orta / Aşağı); şirkət öz kontekstinə görə dəyişməlidir.", False),
    ("'(yoxlanmalı)' qeydi olan dəyərlər mənbədə birbaşa təsdiqlənməyib.", False),
]
for i, (t, bold) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = Font(name=F, bold=bold, size=14 if i == 1 else 10, color="0B3A4A" if bold else "141922")
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 120

# 2. Sources
src_rows = []
for l in laws:
    ref = g(l, "law_ref", "ref") or ", ".join(x for x in [g(l, "number", "act"), g(l, "adopted", "accept_date", "date")] if x)
    full = g(l, "law_full", "full_title_az", "title_az", "title", "official_title")
    mine = [o for o in obls if o.get("law_full") == full or (l.get("key") and str(o.get("id", "")).startswith(l["key"] + "-"))]
    short = g(l, "law_short", "short", "short_name") or next((o["law_short"] for o in mine if o.get("law_short")), full)
    auths = [o.get("authority") for o in mine if o.get("authority")]
    auth = (g(l, "authority", "supervisory_authority") or ("Azərbaycan Respublikasının Mərkəzi Bankı" if "Mərkəzi Bank" in ref else "")
            or (max(set(auths), key=auths.count) if auths else ""))
    if ref.strip("—, ") == "":
        ref = ""
    src_rows.append([short, full,
                     ref, auth, g(l, "status", "status_per_source") or "Qüvvədədir", g(l, "effective_date"),
                     g(l, "version", "last_amendment", "last_amendment_in_force", "last_amended"), g(l, "source_url", "url")])
sheet(wb, "Mənbələr", ["Qısa ad", "Rəsmi ad", "Nömrə və tarix", "Nəzarət orqanı", "Status", "Qüvvəyə minmə", "Versiya / son dəyişiklik", "Mənbə linki"],
      src_rows, [30, 60, 24, 34, 30, 14, 30, 50], "Sources")

# 3. Obligations
orows = []
for o in obls:
    orows.append([o.get("id"), CAT.get(o.get("category"), o.get("category")), o.get("law_short"), o.get("article"),
                  o.get("obligation_az"), o.get("obligation_en"), o.get("verbatim_az"), o.get("applies_to"),
                  o.get("type"), o.get("frequency"), RISK.get(o.get("risk"), o.get("risk")), o.get("authority"),
                  j(o.get("suggested_controls")), j(o.get("related_policies")), o.get("source_url")])
ws = sheet(wb, "Öhdəliklər", ["ID", "Kateqoriya", "Qanun / qayda", "Maddə / bənd", "Öhdəlik (AZ)", "Obligation (EN)",
                               "Rəsmi mətn (sitat)", "Kimə aiddir", "Növ", "Tezlik / son tarix", "Risk", "Nəzarət orqanı",
                               "Təklif olunan control-lar", "Əlaqəli policy-lər", "Mənbə"],
           orows, [8, 16, 26, 12, 40, 40, 70, 22, 11, 22, 10, 26, 40, 32, 40], "Obligations")
for r, o in enumerate(obls, 2):
    f = RISKFILL.get(o.get("risk"))
    if f:
        ws.cell(row=r, column=11).fill = PatternFill("solid", fgColor=f)

# 4. Policies
prow = []
for p in pols:
    drv = " · ".join(x for x in [g(p, "driving_articles", "articles", "driving_paragraphs", "paragraphs"), g(p, "drivers", "driven_by", "obligations")] if x)
    prow.append([g(p, "name_az", "name"), g(p, "name_en"), g(p, "purpose", "purpose_az"), drv, p["_from"]])
sheet(wb, "Policy-lər", ["Policy / prosedur (AZ)", "Policy (EN)", "Məqsəd", "Hansı maddələrdən irəli gəlir", "Qanun"],
      prow, [36, 34, 60, 40, 34], "Policies")

# 5. Summary with formulas
ws = wb.create_sheet("Xülasə")
ws.append(["Qanun / qayda", "Öhdəlik sayı", "Yüksək", "Orta", "Aşağı"])
names = sorted({o.get("law_short") for o in obls if o.get("law_short")})
for i, n in enumerate(names, 2):
    ws.append([n, f'=COUNTIF(Öhdəliklər!$C:$C,A{i})', f'=COUNTIFS(Öhdəliklər!$C:$C,A{i},Öhdəliklər!$K:$K,"Yüksək")',
               f'=COUNTIFS(Öhdəliklər!$C:$C,A{i},Öhdəliklər!$K:$K,"Orta")', f'=COUNTIFS(Öhdəliklər!$C:$C,A{i},Öhdəliklər!$K:$K,"Aşağı")'])
last = len(names) + 1
ws.append(["Cəmi"] + [f"=SUM({c}2:{c}{last})" for c in "BCDE"])
for c in ws[1]:
    c.font, c.fill = HFONT, HEAD
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name=F, size=10, bold=row[0].value == "Cəmi")
ws.column_dimensions["A"].width = 46
for c in "BCDE":
    ws.column_dimensions[c].width = 14

# 6. Notes
ws = wb.create_sheet("Qeydlər")
ws.append(["Fayl", "Qeyd"])
for f, n in notes:
    ws.append([f.replace(".json", ""), n])
for c in ws[1]:
    c.font, c.fill = HFONT, HEAD
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.font, c.alignment = BODY, WRAP
ws.column_dimensions["A"].width = 18
ws.column_dimensions["B"].width = 130

from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
xlsx = OUT / "AZ-oehdelik-kitabxanasi.xlsx"
wb.save(xlsx)

# GRC360 import CSV (columns mirror exportObligationsCsv)
with open(OUT / "GRC360-obligations-import.csv", "w", newline="", encoding="utf-8") as fh:
    w = csv.writer(fh, quoting=csv.QUOTE_ALL)
    w.writerow(["legalActName", "articleOrClauseNo", "category", "requirementDescription", "regulatoryAuthority", "status", "riskLevel"])
    for o in obls:
        name = f'{o.get("law_full") or o.get("law_short")} ({o.get("law_ref")})' if o.get("law_ref") else (o.get("law_full") or o.get("law_short"))
        desc = f'[{o.get("id")}] {o.get("obligation_az")} — Mətn: "{o.get("verbatim_az")}"'
        clean = lambda s: " ".join(str(s or "").split())
        w.writerow([clean(name)[:500], clean(o.get("article"))[:255], o.get("category"), clean(desc),
                    clean(o.get("authority")).replace("(yoxlanmalı)", "").strip()[:255], "UNDER_REVIEW", o.get("risk") or "MEDIUM"])
print(f"{len(laws)} laws, {len(obls)} obligations, {len(pols)} policies -> {OUT}")
