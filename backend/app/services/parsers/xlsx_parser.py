import io

from openpyxl import load_workbook

from app.services.document_model import ParsedDocument, Sheet


def parse_xlsx(filename: str, content: bytes) -> ParsedDocument:
    workbook = load_workbook(io.BytesIO(content), data_only=True)
    sheets: list[Sheet] = []

    for ws in workbook.worksheets:
        all_rows = [
            ["" if cell is None else str(cell) for cell in row]
            for row in ws.iter_rows(values_only=True)
        ]
        if not all_rows:
            sheets.append(Sheet(name=ws.title, columns=[], rows=[]))
            continue
        columns = all_rows[0]
        data_rows = all_rows[1:]
        sheets.append(Sheet(name=ws.title, columns=columns, rows=data_rows))

    return ParsedDocument(doc_type="xlsx", filename=filename, sheets=sheets)
