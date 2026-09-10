import io

from docx import Document as DocxDocument

from app.services.document_model import Block, ParsedDocument

_HEADING_STYLES = {f"Heading {i}": i for i in range(1, 7)}


def parse_docx(filename: str, content: bytes) -> ParsedDocument:
    """Parse a .docx into ordered blocks.

    Paragraphs (including headings) are emitted first, in document order,
    followed by table cells grouped by table/row/col. python-docx does not
    cheaply expose true interleaved body order for typed table objects, and
    MVP scope only requires preserving each paragraph's text and each
    table's own row/column order — not exact paragraph/table interleaving.
    """
    doc = DocxDocument(io.BytesIO(content))
    blocks: list[Block] = []
    index = 0

    for para in doc.paragraphs:
        heading_level = _HEADING_STYLES.get(para.style.name if para.style else "")
        blocks.append(
            Block(
                index=index,
                kind="heading" if heading_level else "paragraph",
                text=para.text,
                heading_level=heading_level,
            )
        )
        index += 1

    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                blocks.append(
                    Block(
                        index=index,
                        kind="table_cell",
                        text=cell.text,
                        table_index=t_idx,
                        row_index=r_idx,
                        col_index=c_idx,
                    )
                )
                index += 1

    return ParsedDocument(doc_type="docx", filename=filename, blocks=blocks)
