from app.services.document_model import ParsedDocument, Sheet
from app.services.dlp.engine import scan_document
from app.services.export.exporter import export_csv, export_xlsx
from app.services.session_store import Session
from app.services.tokenization.sanitizer import sanitize_session
import io
import time

from openpyxl import load_workbook


def _build_document() -> ParsedDocument:
    sheet = Sheet(
        name="Employees",
        columns=["Name", "Salary", "Email", "Department"],
        rows=[
            ["Aydin Huseynov", "5000 USD", "aydin@example.com", "Engineering"],
            ["Leyla Aliyeva", "4200 USD", "leyla@example.com", "Finance"],
            ["Kamran Mammadov", "3900 USD", "kamran@example.com", "Sales"],
        ],
    )
    return ParsedDocument(doc_type="xlsx", filename="employees.xlsx", sheets=[sheet])


def _session_with_scanned_and_approved_document() -> Session:
    document = _build_document()
    session = Session(id="s1", created_at=time.time(), expires_at=time.time() + 3600, ttl_seconds=3600)
    session.document = document
    findings = scan_document(document)
    session.findings = {f.id: f for f in findings}
    session.approved_ids = {f.id for f in findings}
    sanitize_session(session)
    return session


def test_row_order_preserved_after_sanitize_and_export():
    session = _session_with_scanned_and_approved_document()
    original_names = [row[0] for row in session.document.sheets[0].rows]

    csv_bytes = export_csv(session, "rehydrated")
    lines = csv_bytes.decode().strip().splitlines()
    exported_names = [line.split(",")[0] for line in lines[1:]]

    assert exported_names == original_names


def test_column_order_preserved_after_export():
    session = _session_with_scanned_and_approved_document()
    original_columns = session.document.sheets[0].columns

    csv_bytes = export_csv(session, "sanitized")
    header = csv_bytes.decode().strip().splitlines()[0].split(",")

    assert header == original_columns


def test_worksheet_name_preserved_in_xlsx_export():
    session = _session_with_scanned_and_approved_document()
    xlsx_bytes = export_xlsx(session, "rehydrated")
    workbook = load_workbook(io.BytesIO(xlsx_bytes))
    assert workbook.sheetnames == ["Employees"]


def test_unrelated_department_column_untouched():
    session = _session_with_scanned_and_approved_document()
    xlsx_bytes = export_xlsx(session, "rehydrated")
    workbook = load_workbook(io.BytesIO(xlsx_bytes))
    ws = workbook["Employees"]
    departments = [row[3] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert departments == ["Engineering", "Finance", "Sales"]
